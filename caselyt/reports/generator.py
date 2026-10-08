import io
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from documents.models import Case, Document
from evidence.models import Evidence, Finding
from timeline.models import TimelineEvent
from contradictions.models import Contradiction, MissingInformation

class CaseReportGenerator:
    @staticmethod
    def generate_docx_report(case_id: str = None) -> io.BytesIO:
        case_obj = Case.objects.first()
        doc = docx.Document()

        # Title
        p_title = doc.add_paragraph()
        run_title = p_title.add_run("CASELYT — EVIDENCE INTELLIGENCE DOSSIER")
        run_title.font.name = 'Georgia'
        run_title.font.size = Pt(20)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(108, 32, 41) # #6C2029

        p_sub = doc.add_paragraph()
        run_sub = p_sub.add_run(f"Case Reference: {case_obj.case_number if case_obj else 'FIR 204/2024'} | {case_obj.title if case_obj else 'State vs. Vikram Malhotra'}")
        run_sub.font.size = Pt(11)
        run_sub.font.italic = True

        doc.add_heading("1. Case Overview", level=1)
        doc.add_paragraph(case_obj.description if case_obj else "Evidence-first dossier compiled by CASELYT Evidence Intelligence Engine.")

        # Key Facts
        doc.add_heading("2. Verified Key Facts", level=1)
        findings = Finding.objects.all()
        for f in findings:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"[{f.verification_status}] ").bold = True
            p.add_run(f"{f.fact_statement} (Source: {f.document_id_str}, P.{f.page_number}, {f.paragraph_id})")

        # Timeline
        doc.add_heading("3. Chronological Event Timeline", level=1)
        timeline = TimelineEvent.objects.all()
        for t in timeline:
            p = doc.add_paragraph()
            p.add_run(f"{t.date_str} — {t.title}\n").bold = True
            p.add_run(f"Source Text: \"{t.source_text}\" (Ref: {t.document_id_str}, P.{t.page_number})\n")
            if t.has_conflict:
                p.add_run(f"WARNING CONFLICT: {t.conflict_details}").font.color.rgb = RGBColor(161, 109, 43)

        # Contradictions
        doc.add_heading("4. Contradiction Radar Analysis", level=1)
        conflicts = Contradiction.objects.all()
        for c in conflicts:
            p = doc.add_paragraph()
            p.add_run(f"CONFLICT [{c.category}]: {c.title}\n").bold = True
            p.add_run(f"Source A ({c.source_a_doc_name}, P.{c.source_a_page}): \"{c.source_a_text}\"\n")
            p.add_run(f"VS\n").bold = True
            p.add_run(f"Source B ({c.source_b_doc_name}, P.{c.source_b_page}): \"{c.source_b_text}\"\n")

        # Missing Information
        doc.add_heading("5. Missing Information Audit", level=1)
        missing = MissingInformation.objects.all()
        for m in missing:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"[{m.category}] {m.topic}: ").bold = True
            p.add_run(f"{m.reason_flagged}. Impact: {m.impact_assessment}")

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def generate_pdf_report(case_id: str = None) -> io.BytesIO:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'CaseTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#6C2029')
        )

        h2_style = ParagraphStyle(
            'CaseH2',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=13,
            leading=16,
            textColor=colors.HexColor('#171717'),
            spaceBefore=12,
            spaceAfter=6
        )

        body_style = ParagraphStyle(
            'CaseBody',
            parent=styles['BodyText'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor('#171717')
        )

        elements = []
        case_obj = Case.objects.first()

        elements.append(Paragraph("CASELYT — EVIDENCE INTELLIGENCE DOSSIER", title_style))
        elements.append(Paragraph(f"<b>Case File:</b> {case_obj.title if case_obj else 'State vs. Vikram Malhotra'} | <b>Jurisdiction:</b> High Court of Delhi", body_style))
        elements.append(Spacer(1, 10))
        elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#6C2029'), spaceAfter=15))

        # Key Facts Table
        elements.append(Paragraph("VERIFIED KEY FACTS & EVIDENCE", h2_style))
        findings = Finding.objects.all()
        fact_data = [["Status", "Fact Statement", "Source Citation"]]
        for f in findings:
            fact_data.append([
                f.verification_status,
                f.fact_statement[:80] + "...",
                f"{f.document_id_str} (P.{f.page_number})"
            ])

        t_facts = Table(fact_data, colWidths=[80, 320, 140])
        t_facts.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#EAE5D9')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor('#171717')),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,-1), 8),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#AD8748')),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ]))
        elements.append(t_facts)
        elements.append(Spacer(1, 15))

        # Contradictions
        elements.append(Paragraph("CONTRADICTION RADAR DETECTED CONFLICTS", h2_style))
        conflicts = Contradiction.objects.all()
        for c in conflicts:
            conflict_text = f"<b>[{c.category}] {c.title}</b><br/>Source A ({c.source_a_doc_name}): \"{c.source_a_text}\"<br/><b>VS</b><br/>Source B ({c.source_b_doc_name}): \"{c.source_b_text}\""
            elements.append(Paragraph(conflict_text, body_style))
            elements.append(Spacer(1, 6))

        doc.build(elements)
        buffer.seek(0)
        return buffer
